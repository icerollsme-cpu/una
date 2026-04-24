"""
Excel Output Generator for the SEO Crawler Agent.

Produces a multi-sheet workbook:
  Sheet 1  — Dashboard: all pages, color-coded by opportunity score
  Sheet 2  — High Opportunity: pages scored High (filtered for quick action)
  Sheet 3+ — Per-page detail sheets for High-opportunity pages
"""

import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Color palette (matches seo_agent palette) ────────────────────────────────
C_SECTION   = "1F3864"
C_COL_HDR   = "2E75B6"
C_SUB_HDR   = "4472C4"
C_ROW_ALT   = "DEEAF1"
C_TITLE_BAR = "C55A11"
C_REPORT_HDR= "1F3864"
C_HIGH      = "FF4C4C"
C_MED       = "FFC000"
C_LOW       = "70AD47"
C_NONE      = "BFBFBF"
C_NA        = "F2F2F2"
C_WHITE     = "FFFFFF"
C_LIGHT_GRAY= "F2F2F2"
C_KEY_CELL  = "E8F0F8"
C_WARN      = "FFF0F0"
C_OK        = "F0FFF0"

SCORE_COLORS = {
    "High": C_HIGH,
    "Med":  C_MED,
    "Low":  C_LOW,
    "None": C_NONE,
    "N/A":  C_NA,
}

TITLE_MAX = 60
META_MAX  = 160


def _font(size=9, bold=False, color=C_WHITE, italic=False, name="Calibri"):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def _fill(color: str):
    return PatternFill(fill_type="solid", fgColor=color)


def _align(h="left", v="top", wrap=True, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def _thin_border():
    side = Side(style="thin", color="CCCCCC")
    return Border(left=side, right=side, top=side, bottom=side)


class CrawlerOutputGenerator:
    def __init__(self, output_path: str):
        self.output_path = output_path
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    # ── Public API ────────────────────────────────────────────────────────────

    def generate(self, results: list[dict], site_url: str = ""):
        self._build_dashboard(results, site_url)
        self._build_high_opportunities(results)
        for result in results:
            if result.get("opportunity_score") == "High" and not result.get("error"):
                self._build_page_sheet(result)
        self.wb.save(self.output_path)

    # ── Dashboard ─────────────────────────────────────────────────────────────

    def _build_dashboard(self, results: list[dict], site_url: str):
        ws = self.wb.create_sheet("Dashboard", 0)

        # Title bar
        ws.merge_cells("A1:N1")
        c = ws["A1"]
        ts = datetime.now().strftime("%B %d, %Y at %H:%M")
        c.value = f"SEO Crawler Report  ·  {site_url}  ·  Generated {ts}"
        c.font = _font(size=13, bold=True)
        c.fill = _fill(C_REPORT_HDR)
        c.alignment = _align("center", "center", wrap=False)
        ws.row_dimensions[1].height = 26

        # Stats row
        total    = len(results)
        high     = sum(1 for r in results if r.get("opportunity_score") == "High")
        med      = sum(1 for r in results if r.get("opportunity_score") == "Med")
        low      = sum(1 for r in results if r.get("opportunity_score") == "Low")
        none_    = sum(1 for r in results if r.get("opportunity_score") == "None")
        errors   = sum(1 for r in results if r.get("error"))

        ws.merge_cells("A2:N2")
        c = ws["A2"]
        c.value = (
            f"Pages analyzed: {total}    "
            f"High: {high}    Med: {med}    Low: {low}    "
            f"No change needed: {none_}    Errors: {errors}"
        )
        c.font = _font(size=10, bold=True, color="000000")
        c.fill = _fill("EBF3FB")
        c.alignment = _align("center", "center", wrap=False)
        ws.row_dimensions[2].height = 20

        ws.append([""])  # blank row 3

        # Column headers
        headers = [
            "Page URL",
            "Topic",
            "Main Keyword",
            "Page Type",
            "Opportunity",
            "Reasons",
            "Current Title",
            "Title Len",
            "Suggested Title",
            "Sugg. Len",
            "Current Meta",
            "Meta Len",
            "Suggested Meta",
            "Sugg. Len",
        ]
        ws.append(headers)
        hrow = ws.max_row
        ws.row_dimensions[hrow].height = 22
        for col in range(1, len(headers) + 1):
            c = ws.cell(hrow, col)
            c.font = _font(size=9, bold=True)
            c.fill = _fill(C_COL_HDR)
            c.alignment = _align("center", "center", wrap=True)
            c.border = _thin_border()

        # Data rows — sort by opportunity score priority
        _order = {"High": 0, "Med": 1, "Low": 2, "None": 3, "N/A": 4}
        sorted_results = sorted(results, key=lambda r: _order.get(r.get("opportunity_score", "N/A"), 5))

        for idx, r in enumerate(sorted_results):
            score       = r.get("opportunity_score", "N/A")
            reasons     = " | ".join(r.get("opportunity_reasons") or [])
            curr_title  = r.get("current_title", "")
            curr_tlen   = r.get("current_title_length", len(curr_title))
            sugg_title  = r.get("suggested_title", "")
            sugg_tlen   = r.get("suggested_title_length", len(sugg_title))
            curr_meta   = r.get("current_meta", "")
            curr_mlen   = r.get("current_meta_length", len(curr_meta))
            sugg_meta   = r.get("suggested_meta", "")
            sugg_mlen   = r.get("suggested_meta_length", len(sugg_meta))

            if r.get("error"):
                reasons = f"ERROR: {r['error']}"

            row_vals = [
                r.get("url", ""),
                r.get("topic", ""),
                r.get("main_keyword", ""),
                r.get("page_type", ""),
                score,
                reasons,
                curr_title,
                curr_tlen,
                sugg_title,
                sugg_tlen,
                curr_meta,
                curr_mlen,
                sugg_meta,
                sugg_mlen,
            ]
            ws.append(row_vals)
            drow = ws.max_row
            bg = C_ROW_ALT if idx % 2 == 0 else C_WHITE
            ws.row_dimensions[drow].height = 55

            for col in range(1, len(headers) + 1):
                c = ws.cell(drow, col)
                c.font = _font(size=9, color="000000")
                c.fill = _fill(bg)
                c.alignment = _align("left", "top", wrap=True)
                c.border = _thin_border()

            # Opportunity score cell (col 5)
            score_color = SCORE_COLORS.get(score)
            if score_color:
                sc = ws.cell(drow, 5)
                sc.fill = _fill(score_color)
                sc.font = _font(size=9, bold=True, color=C_WHITE if score in ("High", "Med") else "000000")
                sc.alignment = _align("center", "center")

            # Title length cells — warn if over limit
            _length_cell(ws.cell(drow, 8),  curr_tlen, TITLE_MAX)
            _length_cell(ws.cell(drow, 10), sugg_tlen, TITLE_MAX)
            _length_cell(ws.cell(drow, 12), curr_mlen, META_MAX)
            _length_cell(ws.cell(drow, 14), sugg_mlen, META_MAX)

        # Column widths
        widths = [40, 28, 22, 12, 11, 38, 32, 9, 32, 9, 36, 9, 36, 9]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w

        ws.freeze_panes = "A5"

    # ── High Opportunities sheet ──────────────────────────────────────────────

    def _build_high_opportunities(self, results: list[dict]):
        high = [r for r in results if r.get("opportunity_score") == "High" and not r.get("error")]
        if not high:
            return

        ws = self.wb.create_sheet("High Opportunities")

        ws.merge_cells("A1:F1")
        c = ws["A1"]
        c.value = f"High Opportunity Pages  ·  {len(high)} page(s) needing immediate attention"
        c.font = _font(size=12, bold=True)
        c.fill = _fill(C_HIGH)
        c.alignment = _align("center", "center", wrap=False)
        ws.row_dimensions[1].height = 24

        ws.append([""])  # blank row 2

        headers = ["Page URL", "Main Keyword", "Issue(s)", "Suggested Title", "Suggested Meta", "Quick Fix"]
        ws.append(headers)
        hrow = ws.max_row
        ws.row_dimensions[hrow].height = 20
        for col in range(1, len(headers) + 1):
            c = ws.cell(hrow, col)
            c.font = _font(size=9, bold=True)
            c.fill = _fill(C_COL_HDR)
            c.alignment = _align("center", "center", wrap=True)
            c.border = _thin_border()

        for idx, r in enumerate(high):
            reasons = " | ".join(r.get("opportunity_reasons") or [])
            row_vals = [
                r.get("url", ""),
                r.get("main_keyword", ""),
                reasons,
                r.get("suggested_title", ""),
                r.get("suggested_meta", ""),
                r.get("quick_fix", ""),
            ]
            ws.append(row_vals)
            drow = ws.max_row
            bg = C_ROW_ALT if idx % 2 == 0 else C_WHITE
            ws.row_dimensions[drow].height = 60

            for col in range(1, len(headers) + 1):
                c = ws.cell(drow, col)
                c.font = _font(size=9, color="000000")
                c.fill = _fill(bg)
                c.alignment = _align("left", "top", wrap=True)
                c.border = _thin_border()

        widths = [42, 24, 36, 34, 38, 32]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w

        ws.freeze_panes = "A4"

    # ── Per-page detail sheet ─────────────────────────────────────────────────

    def _build_page_sheet(self, result: dict):
        url = result.get("url", "")
        kw  = result.get("main_keyword") or result.get("topic") or "Page"

        raw_name = re.sub(r'[\\/*?:\[\]]', "", kw)[:26].strip() or "Page"
        used = {s.title for s in self.wb.worksheets}
        sheet_name = raw_name
        counter = 2
        while sheet_name in used:
            sheet_name = raw_name[:23] + f" {counter}"
            counter += 1

        ws = self.wb.create_sheet(sheet_name)
        row = 1

        # Title banner
        ws.merge_cells(f"A{row}:D{row}")
        c = ws.cell(row, 1)
        c.value = f"SEO Opportunity: {kw}"
        c.font = _font(size=13, bold=True)
        c.fill = _fill(C_TITLE_BAR)
        c.alignment = _align("center", "center", wrap=False)
        ws.row_dimensions[row].height = 24
        row += 1

        # URL
        ws.merge_cells(f"A{row}:D{row}")
        c = ws.cell(row, 1)
        c.value = url
        c.font = Font(name="Calibri", size=9, italic=True, color="0563C1", underline="single")
        c.alignment = _align("left", "center", wrap=False)
        row += 2

        row = self._section_header(ws, row, "ANALYSIS")
        row = self._kv(ws, row, "Topic",            result.get("topic", ""))
        row = self._kv(ws, row, "Main Keyword",      result.get("main_keyword", ""))
        row = self._kv(ws, row, "Secondary Keywords",", ".join(result.get("secondary_keywords") or []))
        row = self._kv(ws, row, "Page Type",         result.get("page_type", ""))
        score = result.get("opportunity_score", "")
        row = self._kv(ws, row, "Opportunity Score", score,
                       val_color=SCORE_COLORS.get(score))
        reasons = "\n".join(f"• {r}" for r in (result.get("opportunity_reasons") or []))
        row = self._kv(ws, row, "Reasons",           reasons)
        row += 1

        row = self._section_header(ws, row, "TITLE TAG")
        curr_title = result.get("current_title", "")
        sugg_title = result.get("suggested_title", "")
        curr_tlen  = len(curr_title)
        sugg_tlen  = len(sugg_title)
        row = self._kv(ws, row, f"Current Title ({curr_tlen} chars)",  curr_title,
                       val_color=(C_WARN if curr_tlen > TITLE_MAX else None))
        row = self._kv(ws, row, f"Suggested Title ({sugg_tlen} chars)", sugg_title,
                       val_color=C_OK)
        row = self._kv(ws, row, "Change Needed?",
                       "YES" if result.get("title_change_needed") else "No")
        row += 1

        row = self._section_header(ws, row, "META DESCRIPTION")
        curr_meta = result.get("current_meta", "")
        sugg_meta = result.get("suggested_meta", "")
        curr_mlen = len(curr_meta)
        sugg_mlen = len(sugg_meta)
        row = self._kv(ws, row, f"Current Meta ({curr_mlen} chars)",   curr_meta,
                       val_color=(C_WARN if curr_mlen > META_MAX or curr_mlen == 0 else None))
        row = self._kv(ws, row, f"Suggested Meta ({sugg_mlen} chars)", sugg_meta,
                       val_color=C_OK)
        row = self._kv(ws, row, "Change Needed?",
                       "YES" if result.get("meta_change_needed") else "No")
        row += 1

        row = self._section_header(ws, row, "QUICK FIX")
        row = self._kv(ws, row, "Action", result.get("quick_fix", ""))

        for col_letter, width in [("A", 26), ("B", 54), ("C", 12), ("D", 12)]:
            ws.column_dimensions[col_letter].width = width

    # ── Low-level helpers ─────────────────────────────────────────────────────

    def _section_header(self, ws, row: int, title: str, ncols: int = 4) -> int:
        ws.merge_cells(f"A{row}:{get_column_letter(ncols)}{row}")
        c = ws.cell(row, 1)
        c.value = f"  {title}"
        c.font = _font(size=11, bold=True)
        c.fill = _fill(C_SECTION)
        c.alignment = _align("left", "center", wrap=False, indent=1)
        ws.row_dimensions[row].height = 20
        return row + 1

    def _kv(self, ws, row: int, key: str, value, val_color: str = None) -> int:
        kc = ws.cell(row, 1)
        kc.value = key
        kc.font = _font(size=9, bold=True, color="000000")
        kc.fill = _fill(C_KEY_CELL)
        kc.alignment = _align("left", "top", wrap=True)
        kc.border = _thin_border()

        ws.merge_cells(f"B{row}:D{row}")
        vc = ws.cell(row, 2)
        vc.value = str(value) if value is not None else ""
        vc.font = _font(size=9, color="000000")
        vc.fill = _fill(val_color or C_WHITE)
        vc.alignment = _align("left", "top", wrap=True)
        vc.border = _thin_border()

        ws.row_dimensions[row].height = 40
        return row + 1


# ── Utility ───────────────────────────────────────────────────────────────────

def _length_cell(cell, length: int, limit: int):
    """Color a length cell red if over limit, green if OK."""
    if length == 0:
        cell.fill = _fill(C_WARN)
        cell.font = _font(size=9, bold=True, color="CC0000")
    elif length > limit:
        cell.fill = _fill(C_WARN)
        cell.font = _font(size=9, bold=True, color="CC0000")
    else:
        cell.fill = _fill(C_OK)
        cell.font = _font(size=9, color="006400")
    cell.alignment = _align("center", "center", wrap=False)
    cell.border = _thin_border()
